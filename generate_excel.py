from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

data = [
    # 2022
    (date(2022, 10, 11), 140),
    (date(2022, 10, 13), 95),
    (date(2022, 10, 14), 0),
    (date(2022, 10, 17), 140),
    (date(2022, 10, 18), 95),
    (date(2022, 10, 24), 7440),
    (date(2022, 11, 7), 645),
    (date(2022, 11, 22), 450),
    (date(2022, 11, 24), 450),
    (date(2022, 11, 28), 900),
    (date(2022, 11, 30), 165),
    (date(2022, 12, 9), 30),
    (date(2022, 12, 13), 55),
    (date(2022, 12, 14), 2050),
    (date(2022, 12, 15), 90),
    (date(2022, 12, 28), 1400),
    # 2023
    (date(2023, 1, 24), 4220),
    (date(2023, 2, 8), 595),
    (date(2023, 2, 15), 485),
    (date(2023, 2, 16), 500),
    (date(2023, 2, 22), 300),
    (date(2023, 2, 23), 30),
    (date(2023, 3, 1), 250),
    (date(2023, 3, 6), 610),
    (date(2023, 3, 9), 296),
    (date(2023, 3, 14), 100),
    (date(2023, 3, 22), 300),
    (date(2023, 3, 29), 1283),
    (date(2023, 4, 12), 90),
    (date(2023, 4, 13), 1580),
    (date(2023, 5, 22), 600),
    (date(2023, 5, 23), 586),
    (date(2023, 6, 6), 260),
    (date(2023, 6, 9), 500),
    (date(2023, 6, 14), 110),
    (date(2023, 6, 20), 320),
    (date(2023, 7, 18), 205),
    (date(2023, 7, 19), 1060),
    (date(2023, 7, 24), 300),
    (date(2023, 7, 25), 95),
    (date(2023, 9, 23), 10),
    # 2024
    (date(2024, 1, 16), 20),
    (date(2024, 1, 30), 300),
    (date(2024, 7, 4), 180),
    (date(2024, 7, 16), 50),
    (date(2024, 7, 17), 600),
    (date(2024, 7, 18), 90),
    (date(2024, 7, 23), 440),
    (date(2024, 8, 1), 1500),
    (date(2024, 8, 29), 300),
    (date(2024, 9, 10), 50),
    (date(2024, 9, 11), 300),
    (date(2024, 9, 19), 180),
    (date(2024, 9, 23), 1700),
    (date(2024, 9, 30), 40),
    (date(2024, 10, 23), 1000),
    (date(2024, 10, 24), 1250),
    (date(2024, 11, 6), 3000),
    (date(2024, 11, 12), 300),
    (date(2024, 11, 13), 1000),
    (date(2024, 11, 15), 100),
    (date(2024, 11, 19), 10350),
    (date(2024, 11, 21), 1000),
    (date(2024, 11, 27), 300),
    (date(2024, 11, 28), 120),
    (date(2024, 12, 2), 530),
    (date(2024, 12, 4), 650),
    (date(2024, 12, 12), 1420),
    (date(2024, 12, 17), 1773),
    (date(2024, 12, 19), 460),
    (date(2024, 12, 23), 80),
    # 2025
    (date(2025, 1, 9), 350),
    (date(2025, 1, 13), 200),
    (date(2025, 1, 20), 1150),
    (date(2025, 1, 23), 400),
    (date(2025, 1, 27), 175),
    (date(2025, 2, 3), 710),
    (date(2025, 2, 4), 260),
    (date(2025, 2, 12), 400),
    (date(2025, 2, 18), 170),
    (date(2025, 2, 20), 3500),
    (date(2025, 2, 21), 2000),
    (date(2025, 2, 25), 115),
    (date(2025, 2, 27), 2100),
    (date(2025, 3, 3), 250),
    (date(2025, 3, 5), 1150),
    (date(2025, 3, 18), 126),
    (date(2025, 3, 25), 1000),
    (date(2025, 3, 26), 5000),
    (date(2025, 4, 1), 85),
    (date(2025, 4, 8), 300),
    (date(2025, 4, 10), 300),
    (date(2025, 4, 14), 144),
    (date(2025, 4, 21), 2100),
    (date(2025, 4, 22), 200),
    (date(2025, 4, 24), 151),
    (date(2025, 4, 29), 90),
    (date(2025, 4, 30), 200),
    (date(2025, 5, 6), 180),
    (date(2025, 5, 8), 2690),
    (date(2025, 5, 14), 50),
    (date(2025, 5, 22), 621),
    (date(2025, 5, 27), 95),
    (date(2025, 6, 2), 1068),
    (date(2025, 6, 3), 592),
    (date(2025, 6, 4), 400),
    (date(2025, 6, 5), 500),
    (date(2025, 6, 6), 143),
    (date(2025, 6, 11), 150),
    (date(2025, 6, 12), 679),
    (date(2025, 6, 16), 85),
    (date(2025, 6, 17), 140),
    (date(2025, 6, 23), 680),
    (date(2025, 6, 24), 150),
    (date(2025, 6, 25), 250),
    (date(2025, 6, 26), 400),
    (date(2025, 6, 30), 1000),
    (date(2025, 7, 2), 4300),
    (date(2025, 7, 3), 500),
    (date(2025, 7, 7), 1570),
    (date(2025, 7, 9), 400),
    (date(2025, 7, 10), 840),
    (date(2025, 7, 22), 185),
    (date(2025, 7, 23), 190),
    (date(2025, 7, 24), 200),
    (date(2025, 7, 29), 2900),
    (date(2025, 7, 30), 40),
    (date(2025, 7, 31), 200),
    (date(2025, 8, 28), 100),
    (date(2025, 9, 3), 330),
    (date(2025, 9, 4), 90),
    (date(2025, 9, 10), 1215),
    (date(2025, 9, 11), 665),
    (date(2025, 9, 15), 350),
    (date(2025, 9, 16), 300),
    (date(2025, 9, 17), 150),
    (date(2025, 10, 2), 2100),
    (date(2025, 10, 9), 150),
    (date(2025, 10, 13), 90),
    (date(2025, 10, 15), 700),
    (date(2025, 10, 16), 2140),
    (date(2025, 10, 29), 120),
]

data.sort(key=lambda x: x[0])

wb = Workbook()
ws = wb.active
ws.title = "Importes"

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
center = Alignment(horizontal="center")

ws.cell(row=1, column=1, value="Fecha").font = header_font
ws.cell(row=1, column=1).fill = header_fill
ws.cell(row=1, column=1).alignment = center
ws.cell(row=1, column=2, value="Importe").font = header_font
ws.cell(row=1, column=2).fill = header_fill
ws.cell(row=1, column=2).alignment = center

for i, (fecha, importe) in enumerate(data, start=2):
    c1 = ws.cell(row=i, column=1, value=fecha)
    c1.number_format = "DD/MM/YYYY"
    c2 = ws.cell(row=i, column=2, value=importe)
    c2.number_format = '#,##0.00 "€"'

total_row = len(data) + 2
ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
total_cell = ws.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row - 1})")
total_cell.font = Font(bold=True)
total_cell.number_format = '#,##0.00 "€"'

ws.column_dimensions["A"].width = 15
ws.column_dimensions["B"].width = 18

wb.save("/home/user/pictures/importes.xlsx")
print(f"Generated {len(data)} rows")
print(f"Total: {sum(x[1] for x in data)} €")
